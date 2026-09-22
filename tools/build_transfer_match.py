# -*- coding: utf-8 -*-
"""2026학년도 후기 편입생 54명 — 전적대학 ↔ 광주대 계열 매칭표 엑셀 생성.

build_transfer.py 의 명단·전적대학·계열 분류 로직을 그대로 쓰고(§1~§3),
학적 명단에서 영문명·생년월일·성별을 덧붙여 표 하나로 만든다.
"""
import os, warnings, datetime
warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r'C:\projects\tools\build_transfer.py'
code = open(SRC, encoding='utf-8').read().split('# ── 4.')[0]
ns = {}
exec(code, ns)                      # roster / DEPT_GROUP / STAT 확보
roster, DEPT_GROUP, STAT = ns['roster'], ns['DEPT_GROUP'], ns['STAT']

OUT = os.path.join(os.path.dirname(ns['FORM']),
                   '2026학년도 후기 편입생 전적대학·계열 매칭표(광주대학교).xlsx')

# ── 학적 명단에서 영문명·생년월일·성별 보강 ────────────────────────────────
wb = openpyxl.load_workbook(STAT, data_only=True)
info = {}
for r in wb['학위 재학생(신편입 포함) 명단'].iter_rows(min_row=2):
    v = [c.value for c in r]
    if v[5]:
        info[str(v[5])] = dict(영문명=(v[8] or '').strip(), 성별=v[9], 생년월일=v[10])
for s in roster:
    s.update(info.get(str(s['학번']), dict(영문명=None, 성별=None, 생년월일=None)))

# ── 계열 표기(대분류 + 세부) ───────────────────────────────────────────────
GROUP_LABEL = {
    '어문': '인문계열(어문)', '인문': '인문계열(인문)',
    '법정': '사회계열(법정)', '상경': '사회계열(상경)', '사회': '사회계열(사회)',
    '이학': '자연계열(이학)', '공학': '자연계열(공학)',
    '농수산해양기타': '자연계열(농·수산·해양 기타)',
    '의약간호': '의약간호계열', '예체능': '예체능계열', '사범': '사범계열', '통합': '통합계열',
}
MISSING = '자료 없음(확인 필요)'

rows = []
for s in sorted(roster, key=lambda x: (x['학과'], x['학년'], x['이름'])):
    rows.append([
        s['학번'], s['이름'], s['영문명'], s['생년월일'], s['성별'], s['국적'],
        s['학과'], int(s['학년']),
        s['전적대학'] or MISSING, s['전공'] or MISSING,
        GROUP_LABEL[DEPT_GROUP[s['학과']]],
        GROUP_LABEL[s['전적계열']] if s['전적계열'] else MISSING,
    ])

HEAD = ['학번', '한글명', '영문명', '생년월일', '성별', '국적',
        '광주대 학과', '편입학년', '전적대학', '전적대학 학과',
        '광주대 계열', '전적대학 계열']
WIDTH = [12, 16, 30, 12, 6, 12, 22, 9, 26, 30, 20, 20]

# ── 시트 작성 ─────────────────────────────────────────────────────────────
out = openpyxl.Workbook()
ws = out.active
ws.title = '전적대학·계열 매칭표'

NAVY = '1F3864'
ink = Font(name='맑은 고딕', size=10)
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

ws['A1'] = '2026학년도 후기(2026-2학기) 정원외 편입생 전적대학 ↔ 광주대 계열 매칭표'
ws['A1'].font = Font(name='맑은 고딕', size=14, bold=True, color=NAVY)
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
ws['A2'] = (f'등록 54명(3학년 12 · 4학년 42) · 학적 시스템 2026.9.11. 기준 · '
            f'작성 {datetime.date.today():%Y-%m-%d}')
ws['A2'].font = Font(name='맑은 고딕', size=9, color='595959')
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
ws.row_dimensions[1].height = 24

HEAD_ROW = 4
for j, h in enumerate(HEAD, 1):
    c = ws.cell(row=HEAD_ROW, column=j, value=h)
    c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = box
ws.row_dimensions[HEAD_ROW].height = 22

CENTER = {1, 4, 5, 6, 8, 11, 12}
for i, row in enumerate(rows):
    r = HEAD_ROW + 1 + i
    for j, v in enumerate(row, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = ink
        c.border = box
        c.alignment = Alignment(horizontal='center' if j in CENTER else 'left',
                                vertical='center')
        if v == MISSING:
            c.font = Font(name='맑은 고딕', size=10, color='C00000')
            c.fill = PatternFill('solid', fgColor='FCE4E4')
    if i % 2:                                        # 줄무늬로 가독성 확보
        for j in range(1, len(HEAD) + 1):
            if ws.cell(row=r, column=j).value != MISSING:
                ws.cell(row=r, column=j).fill = PatternFill('solid', fgColor='F5F7FA')

last = HEAD_ROW + len(rows)
for j, w in enumerate(WIDTH, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = f'A{HEAD_ROW + 1}'
ws.auto_filter.ref = f'A{HEAD_ROW}:{get_column_letter(len(HEAD))}{last}'

note = last + 2
for k, text in enumerate([
    '※ 광주대 계열·전적대학 계열은 대교협 「서식 2-4 계열별 모집현황」의 계열 구분에 맞춘 것임',
    '※ 전적대학은 전원 외국(중국 등) 대학 — 서식 2-2 출신대학은 외국대학 열에 기재',
    '※ 체육교육(2명)·아동학(1명)은 사범계열, 관광경영은 상경, 국제관계외교는 법정으로 분류함',
    '※ 출처: 학적 = ★20260911 기준 외국인 유학생 재학생 통계 / 전적대학 = 2026.8.26(강향옥).xlsx, 2026-2학기 중국 유학생 지원자 현황(2차).xlsx',
    '※ 다니야노바 굴누르(오시국립대학교)·레티옥(사이공문화예술관광대) 2명은 담당자 확인분(2026.9.22)',
]):
    c = ws.cell(row=note + k, column=1, value=text)
    c.font = Font(name='맑은 고딕', size=9, color='595959')

# ── 요약 시트(서식 2-4 대조용) ────────────────────────────────────────────
ws2 = out.create_sheet('계열 매칭 요약')
from collections import Counter
pairs = Counter((GROUP_LABEL[DEPT_GROUP[s['학과']]],
                 GROUP_LABEL[s['전적계열']] if s['전적계열'] else MISSING) for s in roster)
ws2['A1'] = '광주대 계열 × 전적대학 계열 (서식 2-4 대조용)'
ws2['A1'].font = Font(name='맑은 고딕', size=12, bold=True, color=NAVY)
for j, h in enumerate(['광주대 계열(모집계열)', '전적대학 계열', '인원'], 1):
    c = ws2.cell(row=3, column=j, value=h)
    c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal='center')
    c.border = box
for i, ((a, b), n) in enumerate(sorted(pairs.items()), 1):
    for j, v in enumerate([a, b, n], 1):
        c = ws2.cell(row=3 + i, column=j, value=v)
        c.font = ink
        c.border = box
        c.alignment = Alignment(horizontal='center' if j == 3 else 'left')
tot = ws2.cell(row=4 + len(pairs), column=1, value='계')
tot.font = Font(name='맑은 고딕', size=10, bold=True)
tot.border = box
ws2.cell(row=4 + len(pairs), column=2).border = box
c = ws2.cell(row=4 + len(pairs), column=3, value=sum(pairs.values()))
c.font = Font(name='맑은 고딕', size=10, bold=True)
c.border = box
c.alignment = Alignment(horizontal='center')
for j, w in enumerate([26, 26, 8], 1):
    ws2.column_dimensions[get_column_letter(j)].width = w

out.save(OUT)
print('저장:', OUT)
print('행수', len(rows), '/ 요약 조합', len(pairs), '/ 합계', sum(pairs.values()))
for (a, b), n in sorted(pairs.items()):
    print(f'  {a:22} ← {b:22} {n}')
