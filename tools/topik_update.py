# -*- coding: utf-8 -*-
"""학위과정 재학생 TOPIK 현황 — 학적 조회 자료로 학년 정보를 갱신해 새 파일로 만든다.

- 학적 조회(.xls, cp949)에서 재학생만 가져와 명부를 다시 만들고
- 직전 현황 파일의 TOPIK·트랙·영어성적을 학번으로, 학번이 바뀐 학생은
  이름+생년월일로 이어 붙인다.
- 급수 현황(피벗)은 새 명부로 다시 계산한다.
"""
import collections, io, os, re, shutil, sys
import openpyxl, xlrd
from copy import copy

BASE = r'C:\Users\user\Desktop\★국제협력처\0. 유학생 통계\인증제 TOPIK 관련'
OLD = os.path.join(BASE, '★20261007 학위과정 재학생 TOPIK 현황.xlsx')
QRY = os.path.join(BASE, '외국인(재학,휴학)조회_20261008 174025.xls')
NEW = os.path.join(BASE, '★20261008 학위과정 재학생 TOPIK 현황.xlsx')

COURSE_ORDER = {'학부': 0, '석사': 1, '박사': 2, '교환': 3}
GRADES = ['6급', '5급', '4급', '3급', '2급', '1급', '미취득']


def nows(s):
    return ''.join(str(s or '').split())


def bd6(s):
    """'2005-07-08' -> '050708'"""
    s = str(s or '').strip()
    if '-' in s and len(s) >= 10:
        return s[2:4] + s[5:7] + s[8:10]
    return s


def bd_full(yymmdd, jumin):
    """주민등록번호 뒷자리 첫 숫자로 세기를 판정해 'YYYY-MM-DD'."""
    y = str(yymmdd).strip()
    if len(y) != 6 or not y.isdigit():
        return str(yymmdd or '')
    m = re.search(r'-(\d)', str(jumin or ''))
    c = m.group(1) if m else ''
    century = '19' if c in ('1', '2', '5', '6', '9', '0') else '20'
    return '%s%s-%s-%s' % (century, y[:2], y[2:4], y[4:6])


def read_query():
    sh = xlrd.open_workbook(QRY, encoding_override='cp949').sheet_by_index(0)
    h = sh.row_values(0)
    rows = [dict(zip(h, sh.row_values(i))) for i in range(1, sh.nrows)]
    return [r for r in rows if str(r['학적상태']).strip() == '재학']


def read_old():
    ws = openpyxl.load_workbook(OLD, data_only=True)['재학생 명부']
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)
            if r[5] is not None]
    return head, rows


def classify(r):
    """과정·학적구분."""
    adm, day = str(r['입학구분']).strip(), str(r['입학일'])[:10]
    if adm == '편입교환학생':
        return '교환', '교환학생'
    course = str(r['구분']).strip()
    if day >= '2026-09-01':
        return course, ('편입생' if adm.startswith('편입') else '신입생')
    return course, '재학생'


def main():
    head, old_rows = read_old()
    by_no = {str(d['학번']).strip(): d for d in old_rows}
    by_nm = {}
    for d in old_rows:
        by_nm.setdefault((nows(d['한글명']), bd6(d['생년월일'])), d)
    # 국적별 주 트랙 (신규 학생 트랙 추정용)
    track_by_nat = {}
    nat = collections.defaultdict(collections.Counter)
    for d in old_rows:
        if str(d['학적구분']) != '교환학생':
            nat[str(d['국적'])][str(d['트랙'])] += 1
    for k, v in nat.items():
        track_by_nat[k] = v.most_common(1)[0][0]

    recs, carried_no, carried_nm, fresh = [], 0, 0, []
    for r in read_query():
        no = str(r['학번']).strip()
        o = by_no.get(no)
        if o is not None:
            carried_no += 1
        else:
            o = by_nm.get((nows(r['이름']), str(r['생년월일']).strip()))
            if o is not None:
                carried_nm += 1
        course, status = classify(r)
        birth = o['생년월일'] if o else bd_full(r['생년월일'], r['주민등록번호'])
        d = {
            '과정': course, '학적구분': status, '학년': int(r['학년']) if str(r['학년']).strip() else None,
            '학부(과)': str(r['학부(과)']).strip(), '학번': no, '한글명': str(r['이름']).strip(),
            '영문명': str(r['영문이름']).strip() or (o['영문명'] if o else ''),
            '성별': str(r['성별']).strip(), '생년월일': birth, '국적': str(r['국적']).strip(),
            '입학구분': str(r['입학구분']).strip(), '입학일': str(r['입학일'])[:10],
            '학적상태': str(r['학적상태']).strip(),
        }
        if o:
            for k in ('트랙', 'TOPIK 급수', '급수(숫자)', '회차', '시험일', '총점', '유효기간',
                      '영어성적', '자료 출처'):
                d[k] = o.get(k)
        else:
            fresh.append(d)
            d['트랙'] = track_by_nat.get(d['국적'], '한국어')
            d['TOPIK 급수'] = '미취득'
            for k in ('급수(숫자)', '회차', '시험일', '총점', '유효기간', '영어성적'):
                d[k] = None
            d['자료 출처'] = '학적 조회(신규, 트랙 추정)'
        recs.append(d)

    recs.sort(key=lambda d: (COURSE_ORDER.get(d['과정'], 9), d['학년'] or 99,
                             d['학부(과)'], d['한글명']))
    for i, d in enumerate(recs, 1):
        d['순번'] = i

    shutil.copyfile(OLD, NEW)
    wb = openpyxl.load_workbook(NEW)
    ws = wb['재학생 명부']

    # 본문 다시 쓰기 (모양은 2행 것을 본뜬다)
    style = [copy(c) for c in ws[2]]
    last = ws.max_row
    for i, d in enumerate(recs):
        row = 2 + i
        for j, key in enumerate(head, 1):
            c = ws.cell(row=row, column=j)
            if row > last:
                s = style[j - 1]
                c.font, c.border, c.fill = copy(s.font), copy(s.border), copy(s.fill)
                c.alignment, c.number_format = copy(s.alignment), s.number_format
            c.value = d.get(key)
    end = 1 + len(recs)
    if last > end:
        ws.delete_rows(end + 1, last - end)
    ws.auto_filter.ref = 'A1:%s%d' % (openpyxl.utils.get_column_letter(len(head)), end)

    # 급수 현황 다시 계산
    def bucket(d):
        if d['과정'] == '교환':
            return '교환학생'
        if d['과정'] == '학부':
            return '학부 %d학년' % d['학년']
        return '대학원 ' + d['과정']

    tab = collections.defaultdict(collections.Counter)
    for d in recs:
        g = str(d['TOPIK 급수'] or '미취득').strip() or '미취득'
        tab[bucket(d)][g if g in GRADES else '미취득'] += 1

    ps = wb['TOPIK 급수 현황']
    labels = [ps.cell(row=r, column=1).value for r in range(2, 10)]
    tot = collections.Counter()
    for r, lab in zip(range(2, 10), labels):
        cnt = tot if lab == '계' else tab.get(lab, collections.Counter())
        if lab != '계':
            tot.update(cnt)
        n = sum(cnt.values())
        up4 = sum(cnt[g] for g in ('6급', '5급', '4급'))
        up3 = up4 + cnt['3급']
        vals = [cnt[g] for g in GRADES] + [n, up4, (up4 / n if n else 0), up3,
                                           (up3 / n if n else 0)]
        for j, v in enumerate(vals, 2):
            ps.cell(row=r, column=j).value = v

    ps.cell(row=11, column=1).value = (
        '※ 기준일 : 2026. 10. 8. 학적 조회(재학생 %d명, 휴학 12명 제외) — 학년은 학적시스템 현재 값'
        % len(recs))
    ps.cell(row=12, column=1).value = (
        '※ TOPIK 급수 : 「★20261007 학위과정 재학생 TOPIK 현황」의 급수를 학번(학번 변경자는 '
        '이름·생년월일)으로 이어 붙임 — 학번 일치 %d명, 이름·생년월일 일치 %d명, 신규 %d명'
        % (carried_no, carried_nm, len(fresh)))
    ps.insert_rows(13)
    ps.cell(row=13, column=1).value = (
        '※ 신규 %d명(%s)은 직전 자료에 없어 TOPIK 미취득으로 두었고 트랙은 국적 기준 추정값임 '
        '— 확인 필요' % (len(fresh), ', '.join('%s %s' % (d['한글명'], d['학번']) for d in fresh)))

    wb.save(NEW)
    print('생성:', NEW)
    print('재학생 %d명 (학번일치 %d / 이름일치 %d / 신규 %d)'
          % (len(recs), carried_no, carried_nm, len(fresh)))
    for lab in labels:
        if lab == '계':
            continue
        c = tab.get(lab, collections.Counter())
        n = sum(c.values())
        up3 = sum(c[g] for g in ('6급', '5급', '4급', '3급'))
        print('  %-10s %3d명  3급이상 %3d (%4.1f%%)' % (lab, n, up3, 100 * up3 / n if n else 0))


if __name__ == '__main__':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    main()
